from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[2]
REFERENCE = Path(r"C:\Users\leval\Downloads\управление данными.docx")
CONTENT = ROOT / "reports" / "template-edit" / "content-base.docx"
OUTPUT = ROOT / "reports" / "управление данными.docx"
EXPECTED_SHA = "b607cdac5c4a471a85fb2051812051350fd6960e893cac0f864e6f8af7c5403f"

if sha256(REFERENCE.read_bytes()).hexdigest() != EXPECTED_SHA:
    raise RuntimeError("Исходный документ изменился; нужно повторно проверить образец")

original = Document(REFERENCE)
content = Document(CONTENT)

title_xml = [p._p.xml for p in original.paragraphs[:13]]
heading_proto = deepcopy(original.paragraphs[14]._p)
label_proto = deepcopy(original.paragraphs[16]._p)
body_proto = deepcopy(original.paragraphs[13]._p)
code_proto = deepcopy(original.paragraphs[17]._p)

for paragraph in list(original.paragraphs[13:]):
    paragraph._element.getparent().remove(paragraph._element)


def append_like(prototype, value):
    clone = deepcopy(prototype)
    for child in list(clone):
        if child.tag != qn("w:pPr"):
            clone.remove(child)
    source_run = prototype.find(qn("w:r"))
    run = OxmlElement("w:r")
    if source_run is not None:
        rpr = source_run.find(qn("w:rPr"))
        if rpr is not None:
            run.append(deepcopy(rpr))
    text = OxmlElement("w:t")
    text.text = value
    if value.startswith(" ") or value.endswith(" "):
        text.set(qn("xml:space"), "preserve")
    run.append(text)
    clone.append(run)
    body = original._element.body
    body.insert(len(body) - 1, clone)


heading_map = {
    "1 Назначение работы": "Цель работы:",
    "2.1 Факты отказов диспетчера": "1.1. Витрина фактов отказов диспетчера (загрузка по событию)",
    "2.2 Дневной рейтинг продукции": "1.2. Витрина дневного рейтинга продукции (загрузка по расписанию)",
    "3.1 Запрос 1А Отказы из витрины": "2.1. Число отказов по дням — запрос к витрине",
    "3.2 Запрос 1Б Отказы из оперативных таблиц": "2.2. Число отказов по дням — запрос к оперативной БД",
    "3.3 Запрос 2А Рейтинг из витрины": "2.3. Рейтинг продукции — запрос к витрине",
    "3.4 Запрос 2Б Рейтинг из оперативных таблиц": "2.4. Рейтинг продукции — запрос к оперативной БД",
    "3.5 Результаты выполнения": "2.5. Время и результаты выполнения запросов",
    "4 Итоги и выводы": "Вывод:",
    "Использованные материалы проекта": "Использованные материалы проекта:",
}
skip_headings = {"2 Структура и спецификации витрин", "3 Аналитические запросы и результаты"}

started = False
schema_count = 0
for child in content._element.body.iterchildren():
    if child.tag == qn("w:p"):
        paragraph = Paragraph(child, content._body)
        value = paragraph.text.strip()
        if value == "1 Назначение работы":
            started = True
        if not started or not value or value in skip_headings:
            continue
        if value in heading_map:
            append_like(heading_proto, heading_map[value])
        elif paragraph.style.name == "Code":
            append_like(code_proto, paragraph.text.rstrip())
        else:
            append_like(body_proto, paragraph.text)
    elif child.tag == qn("w:tbl") and started:
        table = Table(child, content._body)
        schema_count += 1
        if schema_count <= 2:
            append_like(label_proto, "Структура таблицы:")
            for row in table.rows[1:]:
                cells = [cell.text for cell in row.cells]
                append_like(code_proto, f"{cells[0]} — {cells[1]} — {cells[2]}")
        else:
            append_like(label_proto, "Результаты выполнения EXPLAIN (ANALYZE, BUFFERS):")
            for row in table.rows[1:]:
                cells = [cell.text for cell in row.cells]
                append_like(body_proto, f"{cells[0]}: витрина — {cells[1]}; оперативная БД — {cells[2]}; ускорение — {cells[3]}; результат — {cells[4]}.")

if [p._p.xml for p in original.paragraphs[:13]] != title_xml:
    raise RuntimeError("Титульный лист изменился")
if len(original.sections) != 2:
    raise RuntimeError("Изменилось число секций")

original.save(OUTPUT)
preserved = OUTPUT.with_name("управление данными.preserved.docx")
with ZipFile(REFERENCE) as source_zip, ZipFile(OUTPUT) as edited_zip, ZipFile(preserved, "w") as final_zip:
    for part in source_zip.infolist():
        data = (edited_zip.read(part.filename)
                if part.filename == "word/document.xml"
                else source_zip.read(part.filename))
        final_zip.writestr(part, data)
preserved.replace(OUTPUT)
print(OUTPUT)
