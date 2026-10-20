from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "Отчет по витринам данных ConditerTrans.docx"
BENCHMARK = (ROOT / "scripts" / "benchmark-data-marts.sql").read_text(encoding="utf-8")


def set_font(style, name, size, bold=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)


def shade(cell, color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tcpr.append(shd)


def add_para(text="", bold_lead=None, style=None):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_table(headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        for column, width in zip(table.columns, widths):
            column.width = Cm(width)
    for i, title in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = title
        shade(cell, "E8EDF2")
        if widths:
            cell.width = Cm(widths[i])
        for run in cell.paragraphs[0].runs:
            run.bold = True
    trpr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    trpr.append(repeat)
    for row_values in rows:
        row = table.add_row()
        for i, value in enumerate(row_values):
            cell = row.cells[i]
            cell.text = str(value)
            if widths:
                cell.width = Cm(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_code(code):
    for line in code.strip().splitlines():
        p = doc.add_paragraph(style="Code")
        p.add_run(line.rstrip())
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def benchmark_query(label, next_label=None):
    marker = re.search(r"(?m)^-- " + re.escape(label) + r"\.", BENCHMARK)
    assert marker, label
    start = BENCHMARK.index("EXPLAIN (ANALYZE, BUFFERS)", marker.start())
    if next_label:
        end_marker = re.search(r"(?m)^-- " + re.escape(next_label) + r"\.", BENCHMARK[start:])
        assert end_marker, next_label
        end = start + end_marker.start()
    else:
        end = len(BENCHMARK)
    return BENCHMARK[start:end].strip()


doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.0)
sec.bottom_margin = Cm(2.0)
sec.left_margin = Cm(2.5)
sec.right_margin = Cm(2.0)

styles = doc.styles
set_font(styles["Normal"], "Times New Roman", 11)
styles["Normal"].paragraph_format.line_spacing = 1.15
styles["Normal"].paragraph_format.space_after = Pt(6)
set_font(styles["Title"], "Times New Roman", 18, True)
styles["Title"].paragraph_format.space_after = Pt(12)
set_font(styles["Heading 1"], "Times New Roman", 14, True)
styles["Heading 1"].paragraph_format.space_before = Pt(13)
styles["Heading 1"].paragraph_format.space_after = Pt(6)
set_font(styles["Heading 2"], "Times New Roman", 12, True)
styles["Heading 2"].paragraph_format.space_before = Pt(10)
styles["Heading 2"].paragraph_format.space_after = Pt(5)
code_style = styles.add_style("Code", 1)
set_font(code_style, "Consolas", 7.7)
code_style.paragraph_format.left_indent = Cm(0.4)
code_style.paragraph_format.space_after = Pt(0)
code_style.paragraph_format.line_spacing = 1.0

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("ConditerTrans · Витрины данных · 2026")
for run in footer.runs:
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(90, 90, 90)

for _ in range(5):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.style = "Title"
p.add_run("Отчёт по витринам данных ConditerTrans")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Структура, обновление и сравнение аналитических запросов")
for _ in range(4):
    doc.add_paragraph()
for label in ["Дисциплина: Управление данными", "Выполнил: ______________________________", "Группа: _________________________________", "Проверил: _______________________________"]:
    p = doc.add_paragraph(label)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for _ in range(5):
    doc.add_paragraph()
p = doc.add_paragraph("2026")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()

doc.add_heading("1 Назначение работы", 1)
add_para("В проекте ConditerTrans используются две аналитические витрины PostgreSQL: факты отказов диспетчера и дневной рейтинг продукции. Цель работы — описать их структуру и обновление, затем сравнить чтение готовых данных с расчётом тех же показателей по оперативным таблицам. Для сравнения использован интервал 30 дней включительно и компания 1b0b8bcf-61c9-44c6-9150-f73da30ce21c.")
add_para("Для наполнения учебной базы добавлены 1200 заказов за 90 дней, назначенных диспетчеру user1@example.com, и 150 фактов отказа. За последние 30 дней в витрине рейтинга получено 77 строк. Четыре плана EXPLAIN ANALYZE предоставлены после выполнения запросов в DataGrip 13 октября 2026 года.")

doc.add_heading("2 Структура и спецификации витрин", 1)
doc.add_heading("2.1 Факты отказов диспетчера", 2)
add_para("Таблица public.dispatcher_rejection_facts хранит один актуальный факт отказа на заказ. Гранулярность — order_id; повторный отказ обновляет ту же строку. Источники: заказ, история изменения статусов, строки заказа, товары и справочник причин отказа.")
add_table(["Поле", "Тип", "Назначение"], [
    ("order_id", "uuid, PK, FK", "Идентификатор заказа; одна строка на заказ"),
    ("company_id", "uuid, FK", "Производственная компания товара"),
    ("rejected_at", "timestamptz", "Момент отказа; день отчёта вычисляется в UTC"),
    ("rejection_reason_id", "uuid, FK", "Причина отказа")
], [3.6, 2.5, 10.4])
add_para("Индекс (company_id, rejected_at) поддерживает выборку по компании и периоду. Внешние ключи к заказу и компании настроены с ON DELETE CASCADE; для причины отказа — RESTRICT. При 150 строках планировщик выбрал последовательное чтение таблицы, что нормально для малого объёма.")
add_para("Фактическая загрузка выполняется при переводе заказа в статус Rejected: OrderRepository.UpdateStatusAsync создаёт запись в order_change_histories и добавляет либо обновляет DispatcherRejectionFact. После сохранения OrderController сбрасывает кэш отчёта и посылает уведомление ReportChanged. Отдельной хранимой процедуры и триггера PostgreSQL в текущей версии нет.")
add_para("Ниже — спецификация эквивалентной хранимой процедуры для событийной загрузки. Это проектный SQL, не установленный миграциями; при его внедрении следует убрать дублирующую запись факта из приложения.")
add_code("""
CREATE OR REPLACE PROCEDURE public.load_dispatcher_rejection_fact(
    p_order_id uuid, p_rejected_at timestamptz, p_reason_id uuid)
LANGUAGE plpgsql AS $$
DECLARE v_company_id uuid; v_company_count integer;
BEGIN
    SELECT (array_agg(DISTINCT p.company_id))[1],
           count(DISTINCT p.company_id)
      INTO v_company_id, v_company_count
    FROM public.order_lines ol
    JOIN public.products p ON p.id = ol.product_id
    WHERE ol.order_id = p_order_id;
    IF v_company_count <> 1 THEN
        RAISE EXCEPTION 'У заказа должна быть одна производственная компания';
    END IF;
    INSERT INTO public.dispatcher_rejection_facts
        (order_id, company_id, rejected_at, rejection_reason_id)
    VALUES (p_order_id, v_company_id, p_rejected_at, p_reason_id)
    ON CONFLICT (order_id) DO UPDATE SET
        company_id = EXCLUDED.company_id,
        rejected_at = EXCLUDED.rejected_at,
        rejection_reason_id = EXCLUDED.rejection_reason_id;
END $$;
""")
add_para("Удаление в действующей системе происходит каскадно при удалении заказа. Если понадобится ограничить срок хранения, проект процедуры удаления по дате выглядит так (не развёрнут; срок хранения сейчас не задан):")
add_code("""
CREATE OR REPLACE PROCEDURE public.delete_dispatcher_rejection_facts_before(
    p_cutoff timestamptz)
LANGUAGE sql AS $$
    DELETE FROM public.dispatcher_rejection_facts
    WHERE rejected_at < p_cutoff;
$$;
""")
add_para("Спецификация событийного триггера: после вставки в историю записи со статусом 4 (Rejected) и причиной отказа вызвать процедуру загрузки. В текущей версии вместо триггера действует описанный выше код приложения; одновременно включать оба механизма не требуется.")
add_code("""
CREATE OR REPLACE FUNCTION public.on_order_rejection_history()
RETURNS trigger LANGUAGE plpgsql AS $$
BEGIN
    IF NEW.order_status = 4 AND NEW.rejection_reason_id IS NOT NULL THEN
        CALL public.load_dispatcher_rejection_fact(
            NEW.order_id, NEW.change_time, NEW.rejection_reason_id);
    END IF;
    RETURN NEW;
END $$;
CREATE TRIGGER trg_order_rejection_history
AFTER INSERT ON public.order_change_histories
FOR EACH ROW EXECUTE FUNCTION public.on_order_rejection_history();
""")
add_para("Задание по расписанию для этой витрины не настроено: загрузка привязана к событию отказа. При необходимости плановое задание могло бы сверять факты с историей, но такого задания в проекте нет.")

doc.add_heading("2.2 Дневной рейтинг продукции", 2)
add_para("Таблица public.dispatcher_product_rating_daily заменила прежнее материализованное представление. Гранулярность строки — дата создания заказа в UTC, производственная компания и название продукта. Товары с одинаковым названием объединяются; показатель order_count считает различные заказы, а не строки заказа и не количество единиц.")
add_table(["Поле", "Тип", "Назначение"], [
    ("order_date", "date, PK", "День создания заказа в UTC"),
    ("company_id", "uuid, PK", "Производственная компания"),
    ("product_name", "varchar(256), PK", "Название продукта"),
    ("order_count", "integer > 0", "Число различных заказов")
], [3.6, 2.5, 10.4])
add_para("Учитываются статусы 1, 2, 3, 5, 6, 7; черновики (0) и отказы (4) исключены. Заказы, содержащие товары разных производственных компаний, также исключаются. Первичный ключ (company_id, order_date, product_name) обслуживает выборку по компании и датам.")
add_para("Загрузка и удаление реализованы одной хранимой функцией из миграции 20261013000000: она блокирует параллельные пересчёты, удаляет прежнее содержимое и вставляет рассчитанный агрегат. DELETE и INSERT выполняются в одной транзакции; до её фиксации читатели видят прежнюю версию таблицы. Отдельной процедуры очистки нет: очистка — первая стадия полной загрузки.")
add_code("""
CREATE FUNCTION public.refresh_dispatcher_product_rating_daily()
RETURNS void LANGUAGE plpgsql AS $function$
BEGIN
    PERFORM pg_advisory_xact_lock(60101320261013::bigint);
    DELETE FROM public.dispatcher_product_rating_daily;
    INSERT INTO public.dispatcher_product_rating_daily
        (order_date, company_id, product_name, order_count)
    SELECT (o.creation_date AT TIME ZONE 'UTC')::date,
           p.company_id, p.name, count(DISTINCT o.id)::integer
    FROM public.orders o
    JOIN public.order_lines ol ON ol.order_id = o.id
    JOIN public.products p ON p.id = ol.product_id
    WHERE o.status IN (1, 2, 3, 5, 6, 7)
      AND NOT EXISTS (
          SELECT 1 FROM public.order_lines other_line
          LEFT JOIN public.products other_product
            ON other_product.id = other_line.product_id
          WHERE other_line.order_id = o.id
            AND (other_product.id IS NULL
                 OR other_product.company_id <> p.company_id)
      )
    GROUP BY (o.creation_date AT TIME ZONE 'UTC')::date,
             p.company_id, p.name;
END;
$function$;
""")
add_para("Спецификация задания по расписанию: Hangfire job dispatcher-product-rating-refresh вызывает указанную функцию каждые пять минут (cron */5 * * * *). Класс ProductRatingMartRefreshJob имеет атрибут DisableConcurrentExecution(300); функция дополнительно сериализует пересчёты через advisory lock. Ручной запуск: SELECT public.refresh_dispatcher_product_rating_daily();")
add_code("""
RecurringJob.AddOrUpdate<ProductRatingMartRefreshJob>(
    "dispatcher-product-rating-refresh",
    job => job.RefreshAsync(CancellationToken.None),
    "*/5 * * * *");
""")
add_para("Событийный триггер для рейтинга в текущей версии не предусмотрен: новые или изменённые заказы отражаются после следующего планового пересчёта. Стоимость пересчёта не включена во время аналитического SELECT.")

doc.add_heading("3 Аналитические запросы и результаты", 1)
add_para("Все запросы используют одну компанию и одинаковые границы периода: от CURRENT_DATE − 29 до начала завтрашнего дня в UTC. Время ниже — поле Execution Time одного плана EXPLAIN (ANALYZE, BUFFERS), без времени планирования и сетевой передачи. Перед запуском в DataGrip выполнено SET TIME ZONE 'UTC'.")
for label, title, nxt in [
    ("1А", "Отказы из витрины", "1Б"),
    ("1Б", "Отказы из оперативных таблиц", "2А"),
    ("2А", "Рейтинг из витрины", "2Б"),
    ("2Б", "Рейтинг из оперативных таблиц", None),
]:
    doc.add_heading(f"3.{['1А','1Б','2А','2Б'].index(label)+1} Запрос {label} {title}", 2)
    add_code(benchmark_query(label, nxt))

doc.add_heading("3.5 Результаты выполнения", 2)
add_table(["Пара", "Витрина", "Оперативная БД", "Ускорение", "Результат"], [
    ("Отказы 1А/1Б", "0,165 мс; 2 hit", "0,816 мс; 177 hit", "4,95 раза", "15 дней; 51 факт за период"),
    ("Рейтинг 2А/2Б", "0,102 мс; 4 hit", "12,150 мс; 19 578 hit", "119,12 раза", "77 строк в каждом плане")
], [2.5, 2.7, 3.2, 2.3, 5.8])
add_para("Для отказов оперативный запрос сортирует 150 событий, выделяет последнее событие на заказ и 51 раз обращается к строкам заказа и товарам. Витрина фильтрует 150 готовых фактов, из которых 51 попадает в период. Результат обеих выборок содержит 15 дат.")
add_para("Для рейтинга оперативный запрос выполняет соединения, проверку состава каждого заказа и группировку; его план содержит 19 578 обращений к страницам в буферном кэше. Витрина читает готовые 77 строк с помощью индекса и обращается к четырём страницам. Shared hit показывает обращения к кэшу, а не физические чтения диска.")
add_para("Планы показывают равное число строк в каждой паре, но не доказывают полного совпадения значений. Для строгой проверки нужно выполнить обе SELECT-части каждой пары без EXPLAIN и сравнить результаты. Это разовые замеры на небольшой учебной базе; для устойчивой оценки следует повторить пары несколько раз после прогрева кэша и сравнить медианы.")

doc.add_heading("4 Итоги и выводы", 1)
add_para("Обе витрины сокращают время чтения аналитических данных на испытанном наборе: примерно в 5 раз для отказов и в 119 раз для рейтинга продукции. Наибольший эффект получен там, где оперативный запрос выполняет соединения и агрегирование большого числа строк, а таблица витрины хранит готовый итог.")
add_para("Плата за быстрый SELECT — загрузка и хранение производных данных. Факты отказов обновляются во время изменения заказа; рейтинг пересчитывается каждые пять минут и потому может отставать от оперативной базы. Полный пересчёт рейтинга и поддержание индексов не включены в приведённые времена чтения.")
add_para("В действующей реализации событийная запись фактов находится в приложении, а не в триггере PostgreSQL. Приведённые в разделе 2.1 SQL-процедуры и триггер — спецификация возможного переноса этой логики в БД; они не применены к рабочей схеме. Действующая хранимая функция рейтинга и задание Hangfire описаны по миграции и коду проекта.")

doc.add_heading("Использованные материалы проекта", 1)
for source in [
    "api/DataAccess/Migrations/20260915000409_AddDispatcherRejectionFacts.cs",
    "api/DataAccess/Migrations/20261013000000_ReplaceProductRatingViewWithTable.cs",
    "api/Infrastructure/Persistence/Orders/OrderRepository.cs",
    "api/API/Hangfire/HangfireServiceCollectionExtensions.cs",
    "api/API/Hangfire/ProductRatingMartRefreshJob.cs",
    "scripts/benchmark-data-marts.sql",
    "Планы EXPLAIN (ANALYZE, BUFFERS), предоставленные при выполнении запросов в DataGrip 13.10.2026"
]:
    add_para(source)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
